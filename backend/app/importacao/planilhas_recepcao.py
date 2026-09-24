"""Leitura das planilhas manuais da recepção (Excel, ago/2021 → 2026).

Cada aba é um mês ("AGOSTO 2021", "MAIO2025", " FEVEREIRO 2026"...). Dentro
dela, os pacientes vêm em blocos separados por linhas de plantão
("PLANTÃO DIURNO - 01/05/2022 - FULANA", "PLANTAO NOTURNO 01/05/2022"...);
cada linha de paciente só tem a HORA. O layout mudou muito com os anos:

  - ago–set/2021: sem CPF;
  - out/2021 → fev/2024: CPF na coluna H, SUS na I (fev/2024: SUS na J);
  - mar/2024 →: data de nascimento na C, CPF na I, SUS na J, endereço, telefone;
  - às vezes a linha está deslocada (CPF ou nome na coluna A).

Por isso nada aqui depende de coluna fixa: nome, CPF, SUS, nascimento e hora
são reconhecidos pelo formato de cada célula. O dia vem da linha de plantão,
com tolerância aos erros de digitação vistos nas planilhas reais
("PLANÃO", "05/08" dentro de setembro, "03/03/2023" em março/2024,
"04/06/205", plantão repetido, plantão sem data, data sozinha numa célula).

A data que vale é a do CALENDÁRIO (os boletins impressos são guardados por
data, não por plantão): quem chegou às 03:00 do plantão noturno de 04/08
entrou em 05/08. A linha de plantão só serve pra saber o dia.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from typing import Iterable, Iterator, Optional

MESES = {"JANEIRO": 1, "FEVEREIRO": 2, "MARCO": 3, "ABRIL": 4, "MAIO": 5, "JUNHO": 6,
         "JULHO": 7, "AGOSTO": 8, "SETEMBRO": 9, "OUTUBRO": 10, "NOVEMBRO": 11, "DEZEMBRO": 12}

# Palavras que aparecem nas colunas de raça/município/observação e nunca são nome
_NAO_NOME = {"SEM", "INFORMADO", "INFORMACAO", "SISTEMA", "PARDA", "BRANCA", "PRETA", "AMARELA",
             "INDIGENA", "EXTREMOZ", "NATAL", "RN", "TROCA", "CEARA", "MIRIM", "NAO", "CPF", "SUS",
             "NOME", "COMPLETO", "IDADE", "SEXO", "RACA", "MUNICIPIO", "RESIDENCIA", "HORA", "BAU",
             "MUDANCA", "ENDERECO", "DATA", "NASCIMENTO", "SAO", "GONCALO", "AMARANTE", "PLANTAO",
             "DE", "DA", "DO", "DOS", "DAS", "E", "N"}  # preposições sozinhas não fazem um nome
_RE_PLANTAO = re.compile(r"PLA[NT]*[AÃ]?O|PLANT|PLATAO|PLANAO")
_RE_DATA = re.compile(r"(\d{1,2})\s*[/.-]\s*(\d{1,2})(?:\s*[/.-]\s*(\d{2,4}))?")
_RE_HORA = re.compile(r"^\s*(\d{1,2})\s*[:;.h]\s*>?\s*(\d{2})")


def sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def nome_busca(nome: str) -> str:
    """Forma usada pra comparar/buscar nome: maiúsculo, sem acento, espaços simples."""
    return " ".join(re.sub(r"[^A-Z ]", " ", sem_acento(nome).upper()).split())


@dataclass
class Entrada:
    data_plantao: date
    turno: str                 # DIURNO | NOTURNO
    hora: Optional[time]
    nome: str
    cpf: Optional[str]         # 11 dígitos (formatado vale mesmo com dígito errado — a planilha manda)
    cns: Optional[str]         # 15 dígitos
    nascimento: Optional[date]
    aba: str
    linha: int
    arquivo: str = ""

    @property
    def data(self) -> date:
        """Dia em que o paciente chegou (calendário): madrugada do noturno = dia seguinte."""
        if self.turno == "NOTURNO" and self.hora is not None and self.hora < time(7, 0):
            return self.data_plantao + timedelta(days=1)
        return self.data_plantao

    @property
    def nome_busca(self) -> str:
        return nome_busca(self.nome)


def mes_da_aba(nome_aba: str) -> tuple[Optional[int], Optional[int]]:
    """'MAIO2025' -> (5, 2025); 'DEZEMBRO' -> (12, None); 'Página42' -> (None, None)."""
    s = sem_acento(nome_aba).upper()
    mes = next((n for nome, n in MESES.items() if nome in s), None)
    ano = re.search(r"(20\d{2})", s)
    return mes, (int(ano.group(1)) if ano else None)


# ── Células ─────────────────────────────────────────────────────────────────
def _texto(v) -> str:
    return "" if v is None else str(v).strip()


def _digitos(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\D", "", str(v)) if v is not None else ""


def _cpf(v) -> Optional[str]:
    if isinstance(v, (datetime, date, time)):
        return None
    d = _digitos(v)
    if len(d) == 10 and isinstance(v, (int, float)):
        d = d.zfill(11)  # zero da frente perdido no Excel
    if len(d) != 11 or d == d[0] * 11:
        return None
    s = _texto(v)
    if isinstance(v, str) and not re.fullmatch(r"[\d.\-\s=]+", s):
        return None  # tem letra/parênteses: telefone, observação...
    formatado = isinstance(v, str) and re.search(r"[.\-]", s)
    # formatado como CPF vale mesmo com erro de digitação; número solto só se o
    # dígito verificador bater (senão pode ser telefone)
    return d if formatado or cpf_valido(d) else None


def cpf_valido(d: str) -> bool:
    if len(d) != 11 or d == d[0] * 11:
        return False
    for n in (9, 10):
        soma = sum(int(d[i]) * (n + 1 - i) for i in range(n))
        if (soma * 10 % 11) % 10 != int(d[n]):
            return False
    return True


def _cns(v) -> Optional[str]:
    if isinstance(v, (datetime, date, time)):
        return None
    d = _digitos(v)
    return d if len(d) == 15 else None


def _hora(v) -> Optional[time]:
    if isinstance(v, time):
        return v.replace(second=0, microsecond=0)
    if isinstance(v, datetime):
        return None  # datetime é nascimento/data, nunca hora
    m = _RE_HORA.match(_texto(v))
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return time(int(m.group(1)), int(m.group(2)))
    return None


def _nome(v) -> Optional[str]:
    s = _texto(v)
    if not s or any(c.isdigit() for c in s) or isinstance(v, (int, float)):
        return None
    palavras = nome_busca(s).split()
    if len(palavras) < 2 or all(p in _NAO_NOME for p in palavras):
        return None
    if palavras[0] in ("R", "RUA", "AV", "TV", "TRAVESSA", "AVENIDA", "RODOVIA", "PO"):
        return None  # endereço
    return " ".join(s.split()).upper()


def _nascimento(v, ano_ref: int) -> Optional[date]:
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date) and 1900 <= v.year <= ano_ref:
        return v
    return None


# ── Linhas ──────────────────────────────────────────────────────────────────
def _dia_do_texto(junto: str, mes: int) -> Optional[int]:
    """Dia escrito na linha de plantão. Prefere uma data cujo mês bata com a aba
    ("05/08" dentro de setembro ainda vale o dia 5, se for a única); entende
    DDMM grudado ("1403/2022")."""
    candidatos = []
    for m in _RE_DATA.finditer(junto):
        d, mm = int(m.group(1)), int(m.group(2))
        if 1 <= d <= 31:
            candidatos.append((d, mm))
    grudado = re.search(r"\b(\d{2})(\d{2})\s*/\s*\d{2,4}", junto)
    if grudado and 1 <= int(grudado.group(1)) <= 31:
        candidatos.insert(0, (int(grudado.group(1)), int(grudado.group(2))))
    if not candidatos:
        return None
    return next((d for d, mm in candidatos if mm == mes), candidatos[0][0])


def _eh_separador(celulas: list, mes: int, ano: int) -> tuple[bool, Optional[str], Optional[int], bool]:
    """(é linha de plantão?, turno ou None, dia ou None, é do mês seguinte?)."""
    textos = [_texto(c) for c in celulas if c not in (None, "")]
    if not textos:
        return False, None, None, False
    junto = sem_acento(" ".join(textos)).upper()
    turno = "NOTURNO" if re.search(r"NOT|NORT", junto) else "DIURNO" if re.search(r"DI?U?R?N|DUIR", junto) else None
    tem_plantao = bool(_RE_PLANTAO.search(junto))
    mes_seguinte = (mes % 12) + 1
    # data em célula de data: só vale se for do mês da aba (nascimento solto não é dia)
    datas = [c for c in celulas if isinstance(c, datetime)]
    do_mes = [c for c in datas if (c.year, c.month) == (ano, mes)]
    dia = do_mes[0].day if do_mes else None
    seguinte = False
    if dia is None and tem_plantao:
        dia = _dia_do_texto(junto, mes)
        # "DIURNO 01/12" no fim de novembro = 1º dia do mês seguinte
        m = _RE_DATA.search(junto)
        seguinte = bool(dia == 1 and m and int(m.group(2)) == mes_seguinte)
    so_data = len(textos) == 1 and bool(do_mes)
    return (tem_plantao or so_data), turno, dia, seguinte


def ler_aba(linhas: Iterable[tuple], aba: str, ano_padrao: Optional[int] = None,
            avisos: Optional[list] = None) -> Iterator[Entrada]:
    """Entradas (pacientes) de uma aba. `linhas`: valores das células por linha
    (openpyxl iter_rows(values_only=True)).

    Em 2 passos: primeiro acha as linhas de plantão e o dia escrito em cada uma;
    depois descarta o dia que contradiz a sequência (erro de digitação: "07/11"
    seguido de "05/11") — mas aceita salto pra frente que continua coerente (a
    aba pode não ter todos os dias: "05/08" → "09/08" → "10/08")."""
    avisos = avisos if avisos is not None else []
    mes, ano = mes_da_aba(aba)
    ano = ano or ano_padrao
    if not mes or not ano:
        return
    inicio = date(ano, mes, 1)
    ultimo_dia = (date(ano + (mes == 12), mes % 12 + 1, 1) - timedelta(days=1)).day

    # 1º passo: classifica as linhas
    itens = []  # (n_linha, celulas, None) paciente | (n_linha, None, [turno, dia]) plantão
    for n_linha, row in enumerate(linhas, start=1):
        celulas = list(row)[:20]
        if not any(c not in (None, "") for c in celulas):
            continue
        sep, turno_sep, dia_sep, seguinte = _eh_separador(celulas, mes, ano)
        tem_nome = any(_nome(c) for c in celulas)
        if sep and not (tem_nome and any(_hora(c) for c in celulas)):
            if seguinte:
                dia_sep = ultimo_dia + 1          # 1º do mês seguinte
            if dia_sep is not None and dia_sep > ultimo_dia + 1:
                dia_sep = None
            itens.append((n_linha, None, [turno_sep, dia_sep]))
        elif tem_nome:
            itens.append((n_linha, celulas, None))

    # 2º passo: dia escrito que quebra a ordem (anterior > ele ou ele > próximo) vira "sem dia"
    seps = [s for _, c, s in itens if s is not None]
    ultimo_bom = 0
    for i, s in enumerate(seps):
        d = s[1]
        if d is None:
            continue
        proximo = next((x[1] for x in seps[i + 1:] if x[1] is not None), None)
        if d < ultimo_bom or (proximo is not None and proximo < d and proximo >= ultimo_bom):
            s[1] = None
            s.append(d)  # guarda o que estava escrito, pro aviso
        else:
            ultimo_bom = d

    # 3º passo: percorre atribuindo dia/plantão
    dia, turno = 1, "DIURNO"   # antes do 1º plantão: diurno do dia 1
    viu = False
    for n_linha, celulas, s in itens:
        if s is not None:
            turno_sep, dia_sep = s[0], s[1]
            anterior_dia, anterior_turno = dia, turno
            if not viu:
                dia = dia_sep if dia_sep and dia_sep <= ultimo_dia else 1
                turno = turno_sep or "DIURNO"
                viu = True
                continue
            if dia_sep is not None:
                dia = dia_sep
            elif (turno_sep or "DIURNO") == "DIURNO" and not (turno_sep is None and anterior_turno == "DIURNO"):
                dia = min(anterior_dia + 1, ultimo_dia + 1)
                if len(s) > 2:
                    avisos.append(f"{aba} linha {n_linha}: dia {s[2]} fora de ordem — considerado "
                                  f"{inicio + timedelta(days=dia - 1):%d/%m/%Y}")
            if turno_sep is None or (turno_sep == anterior_turno and dia == anterior_dia):
                # sem turno escrito, ou o mesmo turno repetido no mesmo dia: é o outro plantão
                turno = "NOTURNO" if (dia == anterior_dia and anterior_turno == "DIURNO") else "DIURNO"
            else:
                turno = turno_sep
            continue

        nomes = [(i, n) for i, n in ((i, _nome(c)) for i, c in enumerate(celulas)) if n]
        nome = next((n for i, n in nomes if i == 1), nomes[0][1])  # coluna B primeiro
        cpf = next((c for c in map(_cpf, celulas) if c), None)
        cns = next((c for c in map(_cns, celulas) if c), None)
        hora = next((h for h in map(_hora, celulas) if h), None)
        nasc = next((d for d in (_nascimento(c, ano) for c in celulas) if d), None)
        yield Entrada(inicio + timedelta(days=dia - 1), turno, hora, nome, cpf, cns, nasc, aba, n_linha)


def ano_das_abas(nomes_abas: list[str]) -> dict[str, Optional[int]]:
    """Abas sem ano ('DEZEMBRO', 'SETEMBRO') herdam o ano da vizinhança: o da
    próxima aba com ano, recuando um se o mês for maior (DEZEMBRO antes de
    JANEIRO 2022 = 2021)."""
    resultado: dict[str, Optional[int]] = {}
    for i, aba in enumerate(nomes_abas):
        mes, ano = mes_da_aba(aba)
        if ano or not mes:
            resultado[aba] = ano
            continue
        vizinhos = [mes_da_aba(a) for a in nomes_abas[i + 1:] + nomes_abas[:i][::-1]]
        ref = next(((m, a) for m, a in vizinhos if a), None)
        resultado[aba] = None if not ref else (ref[1] - 1 if mes > ref[0] and ref[0] <= 6 else ref[1])
    return resultado


def ler_pasta(pasta, avisos: Optional[list] = None) -> tuple[list[Entrada], list[dict]]:
    """Lê todos os .xlsx da pasta. Devolve (entradas sem duplicata, resumo por aba).

    Arquivos "Cópia de ..." só entram nos meses que não existem nos outros
    arquivos (a cópia de jun/2024 tinha uma versão diferente de abril/2024).
    Duplicata = mesmo dia, mesma hora e mesmo nome (a mesma aba aparece em mais
    de um arquivo)."""
    import warnings
    from pathlib import Path

    import openpyxl

    avisos = avisos if avisos is not None else []
    arquivos = sorted(Path(pasta).glob("*.xlsx"), key=lambda f: (f.name.upper().startswith("CÓPIA") or
                                                                 f.name.upper().startswith("COPIA"), f.name))
    vistos: dict[tuple, Entrada] = {}
    meses_originais: set[tuple[int, int]] = set()
    resumo: list[dict] = []
    for arq in arquivos:
        copia = sem_acento(arq.name).upper().startswith("COPIA")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # células com data inválida no Excel (lidas como erro)
            wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
            _ler_arquivo(wb, arq.name, copia, vistos, meses_originais, resumo, avisos)
    return list(vistos.values()), resumo


def _ler_arquivo(wb, nome_arquivo, copia, vistos, meses_originais, resumo, avisos) -> None:
    try:
        anos = ano_das_abas(wb.sheetnames)
        for ws in wb.worksheets:
            mes, _ = mes_da_aba(ws.title)
            ano = anos.get(ws.title)
            if not mes or not ano:
                continue
            if copia and (ano, mes) in meses_originais:
                resumo.append({"arquivo": nome_arquivo, "aba": ws.title, "lidas": 0, "novas": 0, "pulada": True})
                continue
            if not copia:
                meses_originais.add((ano, mes))
            lidas = novas = 0
            for e in ler_aba(ws.iter_rows(values_only=True), ws.title, ano, avisos):
                lidas += 1
                e.arquivo = nome_arquivo
                chave = (e.data, e.hora, e.nome_busca[:30])
                if chave not in vistos:
                    vistos[chave] = e
                    novas += 1
            resumo.append({"arquivo": nome_arquivo, "aba": ws.title, "lidas": lidas, "novas": novas, "pulada": False})
    finally:
        wb.close()
