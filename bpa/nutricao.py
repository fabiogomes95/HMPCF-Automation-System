"""Produção da NUTRIÇÃO a partir da planilha mensal (DADOS NUTRIÇÃO.xlsx).

Substitui o antigo gerar_producao_nutricionistas.py (que perdia paciente com
CPF errado e jogava o 1º paciente de cada dia na nutricionista do dia anterior).

Como a planilha é (uma aba por mês, ex. "AGO-26"):
  coluna A: data do dia / nome da nutricionista / anotações (FDS, FERIADO...)
  coluna B: nome do paciente   C: nascimento   D: CPF

Regras (confirmadas com o usuário em 24/09/2026):
  1. Cada dia começa na linha com a data (célula de data OU texto "20/jul",
     "08/Jjun", "02.03.2026"...). Data de outro mês dentro da aba vira o mesmo
     dia no mês da aba, com aviso.
  2. O dia INTEIRO é da nutricionista escrita nele — inclusive o 1º paciente,
     que fica na linha da data, antes do nome.
  3. Dia com 2+ nomes ou "TODAS NUT": pacientes divididos igualmente entre elas.
  4. Dia sem nome: fica pendente — a tela pergunta de quem é. Paciente acima
     da 1ª data da aba: a tela pergunta o dia.
  5. CPF errado/ausente: procura o paciente no Firebird por nome + nascimento
     (feito pelo BPA local, que tem o cadastro em memória).

Gera UM arquivo BPA-I com as 3 nutricionistas e todos os dias do mês
(BPA_NUTRICAO_<AAAAMM>.txt), igual ao dos médicos. Nunca escreve nos lotes DD-MM-AAAA.txt da
digitação (incidente de 15/07/2026, ver legado/nutricao_antiga/NUTRICIONISTAS_TODO.md).
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime

import openpyxl

CBO_NUTRI = "223710"
CODIGO_NUTRI = "0301010048"

_MESES = {"JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6,
          "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12}
_NOTAS = ("FDS", "FINAL DE S", "FERIADO", "FIM DE SEMANA")
_RE_DIA_TEXTO = re.compile(r"^\s*(\d{1,2})\s*[./\-]")


def sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def normalizar_nome(s) -> str:
    """Nome pra comparar com o Firebird: maiúsculo, sem acento, espaços simples, 30 letras."""
    return " ".join(sem_acento(str(s or "")).upper().split())[:30].strip()


def competencia_da_aba(nome_aba: str) -> tuple[int, int] | None:
    """'AGO-26' -> (2026, 8); 'SET - 2025' -> (2025, 9)."""
    s = sem_acento(nome_aba).upper()
    mes = next((n for abrev, n in _MESES.items() if abrev in s), None)
    ano = re.search(r"(\d{4}|\d{2})\s*$", s.strip())
    if not mes or not ano:
        return None
    a = int(ano.group(1))
    return (a + 2000 if a < 100 else a, mes)


def limpar_cpf(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, (int, float)):
        s = str(int(valor))
    else:
        s = re.sub(r"\D", "", str(valor))
    return s.zfill(11) if len(s) == 10 else s  # 10 dígitos = zero da frente perdido no Excel


def _data_nasc(valor) -> str:
    """Nascimento como no cache do Firebird: DD/MM/AAAA."""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    s = str(valor or "").strip()
    m = re.match(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", s)
    return f"{int(m.group(1)):02d}/{int(m.group(2)):02d}/{m.group(3)}" if m else ""


def listar_abas(conteudo: bytes) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    return [ws.title for ws in wb.worksheets if competencia_da_aba(ws.title)]


def _chave_nome(s: str) -> str:
    """Primeiro nome comparável: sem acento, maiúsculo, sem letra dobrada
    (NAÍLLA / NAILLA / NAILA -> NAILA)."""
    return re.sub(r"(.)\1+", r"\1", sem_acento(s).upper())


def _distancia(a: str, b: str) -> int:
    """Quantas letras trocar/incluir/tirar pra ir de a até b (Levenshtein)."""
    anterior = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        atual = [i]
        for j, cb in enumerate(b, 1):
            atual.append(min(anterior[j] + 1, atual[j - 1] + 1, anterior[j - 1] + (ca != cb)))
        anterior = atual
    return anterior[-1]


def _qual_nutri(texto: str, nutricionistas: list[dict]) -> list[dict] | None:
    """Nutricionista(s) citada(s) na célula da coluna A; 'TODAS' = todas.

    O 1º nome do cadastro tem que ser IGUAL à palavra da planilha (ignorando
    acento e letra dobrada). Só se nenhuma bater exatamente, vale o começo do
    nome (4+ letras) ou um nome quase igual (erro de digitação) — e
    só quando aponta pra UMA nutricionista, pra não juntar MARIANE com uma
    MARIA que também tenha o CBO de nutrição."""
    s = sem_acento(texto).upper()
    if "TODAS" in s:
        return list(nutricionistas)  # ler_aba troca depois pelas que aparecem na aba
    primeiros = [(n, _chave_nome(n["nome"].split()[0])) for n in nutricionistas if n["nome"].split()]
    achadas = []
    for palavra in re.sub(r"[^A-Z ]", " ", s).split():
        chave = _chave_nome(palavra)
        if len(chave) < 3:
            continue
        iguais = [n for n, p in primeiros if p == chave]
        if not iguais and len(chave) >= 4:
            # erro de digitação ("Barabara", "Mariana"): no máx. 1 letra (2 em nome longo)
            limite = 2 if len(chave) >= 8 else 1
            parecidas = [n for n, p in primeiros if _distancia(chave, p) <= limite]
            if not parecidas:  # nome abreviado ("MARI")
                parecidas = [n for n, p in primeiros if p.startswith(chave)]
            iguais = parecidas if len(parecidas) == 1 else []  # só se apontar pra UMA
        for n in iguais:
            if n not in achadas:
                achadas.append(n)
    return achadas or None


def ler_aba(conteudo: bytes, aba: str, nutricionistas: list[dict]) -> dict:
    """Lê uma aba e devolve os dias com nutricionista(s) e pacientes (ainda sem
    cruzar com o Firebird). `nutricionistas`: [{"nome", "cns"}] do CADMED (CBO 223710)."""
    comp = competencia_da_aba(aba)
    if not comp:
        raise ValueError(f"Não entendi o mês da aba '{aba}' (esperado algo como AGO-26).")
    ano, mes = comp
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ws = wb[aba]

    dias: list[dict] = []
    avisos: list[str] = []
    dia = None
    comecou = False
    for n_linha, row in enumerate(ws.iter_rows(values_only=True), start=1):
        a, nome, nasc, cpf = (tuple(row) + (None,) * 4)[:4]
        if not comecou:  # pula o título até o cabeçalho (coluna A "DATA"; B varia: "NOME:", "DADOS DOS PACIENTES"...)
            comecou = str(a or "").strip().upper() == "DATA" or str(nome or "").strip().upper().startswith("NOME")
            continue

        # ── coluna A: começo de dia, nutricionista ou anotação
        texto_a = "" if a is None or isinstance(a, (datetime, date)) else str(a).strip()
        nova_data = None
        if isinstance(a, (datetime, date)):
            d = a.date() if isinstance(a, datetime) else a
            nova_data = d
            if (d.year, d.month) != (ano, mes):
                try:
                    nova_data = date(ano, mes, d.day)
                    avisos.append(f"Linha {n_linha}: data {d:%d/%m/%Y} fora do mês da aba — considerada {nova_data:%d/%m/%Y}.")
                except ValueError:
                    nova_data = None
                    avisos.append(f"Linha {n_linha}: data {d:%d/%m/%Y} fora do mês da aba e sem dia correspondente — ignorada.")
        elif texto_a and _RE_DIA_TEXTO.match(texto_a):
            try:
                nova_data = date(ano, mes, int(_RE_DIA_TEXTO.match(texto_a).group(1)))
            except ValueError:
                avisos.append(f"Linha {n_linha}: '{texto_a}' não é um dia válido de {mes:02d}/{ano}.")
        if nova_data:
            dia = next((x for x in dias if x["data"] == nova_data), None)
            if dia is None:
                dia = {"data": nova_data, "nutricionistas": [], "pacientes": []}
                dias.append(dia)

        if texto_a and dia is not None:
            nutris = _qual_nutri(texto_a, nutricionistas)
            if nutris and len(nutris) > 1 and "TODAS" not in sem_acento(texto_a).upper() and len(texto_a.split()) == 1:
                avisos.append(f"Linha {n_linha}: '{texto_a}' serve para mais de uma nutricionista cadastrada "
                              f"({', '.join(n['nome'] for n in nutris)}) — confira o dia.")
            if nutris and len(nutris) == 1 and not any(
                _chave_nome(w) == _chave_nome(nutris[0]["nome"].split()[0]) for w in re.sub(r"[^A-Z ]", " ", sem_acento(texto_a).upper()).split()
            ) and "TODAS" not in sem_acento(texto_a).upper():
                avisos.append(f"Linha {n_linha}: '{texto_a}' entendido como {nutris[0]['nome']} — confira.")
            if nutris and "TODAS" in sem_acento(texto_a).upper():
                dia["todas"] = True
            if nutris:
                for n in nutris:
                    if n not in dia["nutricionistas"]:
                        dia["nutricionistas"].append(n)
            elif not nova_data and not any(sem_acento(texto_a).upper().startswith(x) for x in _NOTAS):
                avisos.append(f"Linha {n_linha}: '{texto_a}' na coluna A não é data, nutricionista nem anotação — ignorado.")

        # ── paciente
        if nome and str(nome).strip():
            if dia is None:  # antes da 1ª data da aba: a tela pergunta o dia
                dia = {"data": None, "nutricionistas": [], "pacientes": []}
                dias.append(dia)
                avisos.append(f"Linha {n_linha}: paciente antes da primeira data da aba — escolha o dia.")
            dia["pacientes"].append({
                "linha": n_linha,
                "nome": " ".join(str(nome).split()).upper(),
                "nascimento": _data_nasc(nasc),
                "cpf_planilha": str(cpf or "").strip(),
                "cpf": limpar_cpf(cpf),
            })

    # "TODAS NUT" = as nutricionistas que trabalham no mês (citadas na aba),
    # não todo mundo com o CBO de nutrição no cadastro.
    citadas = [n for n in nutricionistas if any(n in d["nutricionistas"] for d in dias if not d.get("todas"))]
    for d in dias:
        if d.pop("todas", False) and citadas:
            d["nutricionistas"] = list(citadas)

    dias = [d for d in dias if d["pacientes"]]
    dias.sort(key=lambda d: (d["data"] is not None, d["data"] or date.min))
    return {"aba": aba, "competencia": f"{ano}{mes:02d}", "dias": dias, "avisos": avisos}


def dividir(pacientes: list, nutris: list) -> dict[str, list]:
    """Divide os pacientes do dia igualmente (em ordem) entre as nutricionistas."""
    por = {n["cns"]: [] for n in nutris}
    for i, p in enumerate(pacientes):
        por[nutris[i % len(nutris)]["cns"]].append(p)
    return por
