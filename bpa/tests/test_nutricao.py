"""Nutrição: leitura da planilha (bpa/nutricao.py) e geração do arquivo único
(bpa_local/services/nutricao.py), com Firebird falso."""
import base64
import io
from datetime import datetime

import openpyxl
import pytest

import bpa_gerador as bpa
import nutricao as planilha
from bpa_local.cache import cache
from bpa_local.services import nutricao

NUTRIS = [
    ("700000000000001", "MARIANE SOUZA LIMA"),
    ("700000000000002", "BARBARA COSTA"),
    ("700000000000003", "NAILLA PEREIRA"),
]


def _xlsx(linhas, aba="AGO-26") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = aba
    ws.append([None, "DADOS DOS PACIENTES", None, None])
    ws.append(["DATA", "NOME:", "DATA NTO:", "CPF"])
    for linha in linhas:
        ws.append(list(linha))
    wb.create_sheet("Planilha1")  # aba que não é mês: ignorada
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def d(dia, mes=8):
    return datetime(2026, mes, dia)


PLANILHA = [
    ("", "ANTES DA DATA", datetime(1950, 1, 1), "111.444.777-35"),
    (d(1), "PRIMEIRO DO DIA", "01/02/1960", "529.982.247-25"),     # 1º paciente ANTES do nome
    ("BARBARA", "SEGUNDO DO DIA", "02/02/1960", "123.456.789-09"),  # CPF errado → pelo nome
    ("FDS", "SEM CPF NO FIREBIRD", "03/03/1970", ""),
    (d(2), "DIA DOIS A", "01/01/1980", "529.982.247-25"),
    ("Mariane", "DIA DOIS B", "01/01/1980", "52998224725"),
    ("NAÍLLA", "DIA DOIS C", "01/01/1980", "52998224725"),
    (d(3), "DIA TRES", "01/01/1990", "529.982.247-25"),              # sem nome: pendente
    ("20/ago", "DIA VINTE", "", "529.982.247-25"),                    # data em texto
    ("TODAS NUT", "DIA VINTE B", "", "529.982.247-25"),
    (d(5, 7), "DATA DE OUTRO MES", "", "529.982.247-25"),            # 05/07 na aba de agosto
    ("NAILA", "", "", ""),
    (None, "NINGUEM CONHECE", "01/01/2000", "000.000.000-00"),
]

CACHE = [
    {"sus": "", "nome": "PRIMEIRO DO DIA", "dtnasc": "01/02/1960", "cpf": "52998224725", "id": 1},
    {"sus": "", "nome": "SEGUNDO DO DIA", "dtnasc": "02/02/1960", "cpf": "11144477735", "id": 2},
    {"sus": "", "nome": "SEM CPF NO FIREBIRD", "dtnasc": "03/03/1970", "cpf": "", "id": 3},
]


def _nutris():
    return [{"cns": c, "nome": n} for c, n in NUTRIS]


def test_regras_da_planilha():
    r = planilha.ler_aba(_xlsx(PLANILHA), "AGO-26", _nutris())
    assert r["competencia"] == "202608"
    por_data = {x["data"]: x for x in r["dias"]}
    semdia = r["dias"][0]
    assert semdia["data"] is None and [p["nome"] for p in semdia["pacientes"]] == ["ANTES DA DATA"]

    dia1 = por_data[d(1).date()]
    assert [n["nome"] for n in dia1["nutricionistas"]] == ["BARBARA COSTA"]
    assert [p["nome"] for p in dia1["pacientes"]] == ["PRIMEIRO DO DIA", "SEGUNDO DO DIA", "SEM CPF NO FIREBIRD"]

    assert len(por_data[d(2).date()]["nutricionistas"]) == 2                    # MARIANE + NAÍLLA
    assert por_data[d(3).date()]["nutricionistas"] == []                        # pendente
    assert len(por_data[d(20).date()]["nutricionistas"]) == 3                   # TODAS NUT
    assert [n["nome"] for n in por_data[d(5).date()]["nutricionistas"]] == ["NAILLA PEREIRA"]
    assert any("05/07/2026" in a for a in r["avisos"])


def test_limpar_cpf():
    assert planilha.limpar_cpf("529.982.247-25") == "52998224725"
    assert planilha.limpar_cpf("999.999.999-=99") == "99999999999"
    assert planilha.limpar_cpf(1234567890) == "01234567890"  # zero perdido no Excel
    assert planilha.limpar_cpf(None) == ""


def test_dividir_em_ordem():
    n = [{"cns": "A"}, {"cns": "B"}, {"cns": "C"}]
    assert planilha.dividir([1, 2, 3, 4, 5], n) == {"A": [1, 4], "B": [2, 5], "C": [3]}


class _Cur:
    def __init__(self, s_prd):
        self.s_prd = s_prd
        self.res = []

    def execute(self, sql, params=()):
        if "CADMED" in sql:
            self.res = list(NUTRIS)
        elif "S_PRD" in sql:
            self.res = [(dt, cbo) for cns, dt, cbo in self.s_prd if cns == params[0]]

    def fetchall(self):
        return self.res


class _Con:
    def __init__(self, s_prd=()):
        self.s_prd = list(s_prd)

    def cursor(self):
        return _Cur(self.s_prd)

    def close(self):
        pass


@pytest.fixture
def firebird(monkeypatch, tmp_path):
    monkeypatch.setattr(cache, "pacientes", CACHE)
    monkeypatch.setattr(bpa, "BPA_LOTES_DIR", str(tmp_path))
    estado = {"s_prd": []}
    monkeypatch.setattr(bpa, "conectar", lambda: _Con(estado["s_prd"]))
    monkeypatch.setattr(bpa, "buscar_pacientes", lambda con, docs: (
        [{"doc": x, "sem_doc": x.startswith("ID:")} for x in docs], [], []))
    capturado = []

    def montar(pacs, proc, cbo, cns, data, comp, folha, seq):
        capturado.append((cns, data, folha, seq, [p["doc"] for p in pacs], proc, cbo))
        return [f"{cns}|{data}|{p['doc']}" for p in pacs], folha
    monkeypatch.setattr(bpa, "montar_linhas", montar)
    monkeypatch.setattr(bpa, "montar_cabecalho", lambda comp, n, f, linhas: f"CAB|{comp}|{n}|{f}")
    estado["montar"] = capturado
    estado["pasta"] = tmp_path
    return estado


def _b64(conteudo):
    return base64.b64encode(conteudo).decode()


def test_ler_so_abas(firebird):
    r = nutricao.ler({"arquivo_b64": _b64(_xlsx(PLANILHA))})
    assert r == {"ok": True, "abas": ["AGO-26"]}


def test_ler_acha_pacientes(firebird):
    r = nutricao.ler({"arquivo_b64": _b64(_xlsx(PLANILHA)), "aba": "AGO-26"})
    assert r["ok"], r
    dia1 = next(x for x in r["dias"] if x["data"] == "01/08/2026")
    sit = {p["nome"]: (p["situacao"], p["doc"]) for p in dia1["pacientes"]}
    assert sit["PRIMEIRO DO DIA"] == ("ok", "52998224725")
    assert sit["SEGUNDO DO DIA"] == ("corrigido", "11144477735")
    assert sit["SEM CPF NO FIREBIRD"] == ("sem_doc", "ID:3")
    assert r["contagem"]["nao_achado"] >= 1
    assert [n["curto"] for n in r["nutricionistas"]] == ["MARIANE", "BARBARA", "NAILLA"]


def test_gerar_um_arquivo_com_folha_continua(firebird):
    mariane, barbara = NUTRIS[0][0], NUTRIS[1][0]
    firebird["s_prd"] = [(mariane.zfill(15), "20260810", "223710")] * 100  # outro dia já importado
    firebird["s_prd"] += [(barbara.zfill(15), "20260801", "223710")] * 2   # já importado NESTE dia
    # mesmo CNS atendendo como médico no mesmo dia: não é nutrição repetida
    firebird["s_prd"] += [(mariane.zfill(15), "20260802", "225125")] * 3
    r = nutricao.gerar({"competencia": "202608", "dias": [
        {"data": "01/08/2026", "nutricionistas": [barbara], "docs": ["52998224725", "ID:3"]},
        {"data": "02/08/2026", "nutricionistas": [mariane, barbara], "docs": ["A", "B", "C"]},
        {"data": "03/08/2026", "nutricionistas": [mariane], "docs": ["D"]},
    ]})
    assert r["ok"], r
    assert r["arquivo"] == "BPA_NUTRICAO_202608.txt"
    assert r["registros"] == 6
    texto = (firebird["pasta"] / r["arquivo"]).read_bytes().decode("latin-1")
    assert texto.startswith("CAB|202608|6|") and texto.count("\r\n") == 7

    chamadas = {(c[0], c[1]): c for c in firebird["montar"]}
    # Mariane: 100 de nutrição em outro dia + 3 como médico → folha 2, seq 5; dia 3 continua (seq 7)
    assert chamadas[(mariane.zfill(15), "20260802")][2:5] == (2, 5, ["A", "C"])
    assert chamadas[(mariane.zfill(15), "20260803")][2:5] == (2, 7, ["D"])
    # Bárbara: os 2 do próprio dia 01 não contam pra folha, mas avisam
    assert chamadas[(barbara.zfill(15), "20260801")][2:5] == (1, 1, ["52998224725", "ID:3"])
    assert chamadas[(barbara.zfill(15), "20260802")][2:5] == (1, 3, ["B"])
    assert all(c[5] == "0301010048" and c[6] == "223710" for c in firebird["montar"])
    avisos = {n["nome"]: n["ja_importados"] for n in r["nutricionistas"]}
    assert avisos == {"BARBARA COSTA": 2, "MARIANE SOUZA LIMA": 0}
    # nunca escreve lote de digitação
    assert [p.name for p in firebird["pasta"].iterdir()] == ["BPA_NUTRICAO_202608.txt"]


def test_gerar_recusa_dia_sem_nutri_e_fora_do_mes(firebird):
    r = nutricao.gerar({"competencia": "202608", "dias": [{"data": "03/08/2026", "nutricionistas": [], "docs": ["X"]}]})
    assert not r["ok"] and "nutricionista" in r["erro"]
    r = nutricao.gerar({"competencia": "202608", "dias": [{"data": "03/07/2026", "nutricionistas": [NUTRIS[0][0]], "docs": ["X"]}]})
    assert not r["ok"] and "mês" in r["erro"]


def test_cabecalho_com_outro_texto_na_coluna_b():
    """Set/2026: a planilha trocou "NOME:" por "DADOS DOS PACIENTES" no cabeçalho."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SET - 26"
    ws.append([None, "DADOS DOS PACIENTES JANEIRO DE 2026", None, None])
    ws.append(["DATA", "DADOS DOS PACIENTES ", "DATA NTO:", "CPF"])
    ws.append([datetime(2026, 9, 1), "FULANO", datetime(1980, 1, 1), "529.982.247-25"])
    ws.append(["BARBARA", "CICLANO", datetime(1981, 1, 1), "111.444.777-35"])
    buf = io.BytesIO()
    wb.save(buf)
    r = planilha.ler_aba(buf.getvalue(), "SET - 26", _nutris())
    assert [p["nome"] for d in r["dias"] for p in d["pacientes"]] == ["FULANO", "CICLANO"]
