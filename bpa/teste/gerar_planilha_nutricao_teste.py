"""Gera teste/nutricao_teste.xlsx: planilha da nutrição no MESMO formato da
real (DADOS NUTRIÇÃO.xlsx), mas só com os pacientes falsos de
pacientes_falsos.json ("TESTE ..."), pra testar a aba Nutrição do mês num
BPA de teste. Nada de paciente real aqui — o repositório é público.

Situações da planilha real que ela reproduz (aba AGO-26):
  - paciente acima da 1ª data (escolher o dia na tela)
  - 1º paciente do dia na linha da data, ANTES do nome da nutricionista
  - nomes com grafias diferentes (Mariane, NAILA, NAÍLLA) e anotações (FDS, FERIADO)
  - dia com 2 nutricionistas e dia "TODAS NUT" (divididos igualmente)
  - dia sem nome (escolher de quem é) e data escrita como texto ("20/ago")
  - data de outro mês dentro da aba (12/07 na aba de agosto)
  - CPF em formatos bagunçados, CPF com dígito errado (achado pelo nome),
    paciente sem documento (TESTE SEMDOC) e paciente que não existe no Firebird

Uso: bpa\\.venv\\Scripts\\python bpa\\teste\\gerar_planilha_nutricao_teste.py
"""
import json
import random
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

PASTA = Path(__file__).resolve().parent
random.seed(26)

dados = json.loads((PASTA / "pacientes_falsos.json").read_text(encoding="utf-8"))["pacientes"]


def valido(cpf: str) -> bool:
    if len(cpf) != 11 or cpf == cpf[0] * 11:
        return False
    for n in (9, 10):
        soma = sum(int(cpf[i]) * (n + 1 - i) for i in range(n))
        if (soma * 10 % 11) % 10 != int(cpf[n]):
            return False
    return True


com_cpf = [p for p in dados if valido(p["num_cpf"])]
sem_doc = [p for p in dados if not p["num_cpf"] and "SEMDOC" in p["nome"]]
fila = iter(random.sample(com_cpf, len(com_cpf)))


def nasc(p):
    d = p["dtnasc"]
    return datetime(int(d[:4]), int(d[4:6]), int(d[6:]))


def cpf_bagunçado(cpf: str, jeito: int) -> str:
    """Os formatos que aparecem na planilha real."""
    return [
        f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}",
        f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}.{cpf[9:]}",
        cpf,
        f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-={cpf[9:]}",
    ][jeito % 4]


def pac(jeito=0, p=None):
    p = p or next(fila)
    return (p["nome"], nasc(p), cpf_bagunçado(p["num_cpf"], jeito))


def errado(p=None):
    """CPF com o último dígito trocado: o BPA acha pelo nome + nascimento."""
    p = p or next(fila)
    cpf = p["num_cpf"][:10] + str((int(p["num_cpf"][10]) + 3) % 10)
    return (p["nome"], nasc(p), cpf_bagunçado(cpf, 0))


def semdoc(i):
    p = sem_doc[i]
    return (p["nome"], nasc(p), "")


def d(dia, mes=8):
    return datetime(2026, mes, dia)


# (coluna A, paciente) — paciente = (nome, nascimento, cpf) ou None
LINHAS = [
    (None, pac(0)),                              # acima da 1ª data → escolher o dia
    (d(3), pac(0)),                              # 1º do dia ANTES do nome
    ("BARBARA", pac(1)),
    ("FDS", pac(2)),
    (None, errado()),                            # CPF errado → corrigido pelo nome
    (d(4), pac(0)),
    ("Mariane", pac(3)),
    (None, semdoc(0)),                           # sem documento
    (None, pac(1)),
    (d(5), pac(0)),
    ("MARIANE", pac(2)),
    ("NAILA", pac(0)),                           # 2 nomes → divide
    (None, pac(1)),
    (None, pac(2)),
    (None, pac(3)),
    (d(6), pac(0)),                              # sem nome → escolher na tela
    (None, pac(1)),
    (d(7), pac(0)),
    ("TODAS NUT", pac(1)),                       # divide entre as 3
    (None, pac(2)),
    (None, pac(3)),
    (None, pac(0)),
    (None, semdoc(1)),
    (d(10), pac(0)),
    ("NAÍLLA", pac(1)),
    ("FERIADO", errado()),
    (None, ("TESTE PACIENTE QUE NAO EXISTE", datetime(1970, 1, 1), "000.000.000-00")),  # não achado
    (d(12, 7), pac(0)),                          # 12/07 na aba de agosto → vira 12/08, com aviso
    ("BARBARA", pac(1)),
    ("20/ago", pac(2)),                          # data como texto
    ("MARIANE", pac(3)),
    (None, pac(0)),
    (d(21), pac(1)),
    ("BARBARA", semdoc(2)),
    (None, pac(2)),
]


def main() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AGO-26"
    ws.append([None, "DADOS DOS PACIENTES AGOSTO DE 2026 (TESTE)", None, None, None, None])
    ws.append(["DATA", "NOME:", "DATA NTO:", "CPF", "ENDEREÇO", "BAIRRO"])
    # Com um nome (ex. FABIO): toda nutricionista da planilha vira ela -- pra
    # BPA de teste com uma nutricionista só ("TODAS NUT" continua igual).
    so_uma = sys.argv[1].upper() if len(sys.argv) > 1 else ""
    for a, p in LINHAS:
        nome, dn, cpf = p
        if so_uma and a in ("BARBARA", "Mariane", "MARIANE", "NAILA", "NAÍLLA"):
            a = so_uma
        ws.append([a, nome, dn, cpf, "RUA DE TESTE", "CENTRO"])
    for linha in ws.iter_rows(min_row=3):
        for c in (linha[0], linha[2]):
            if isinstance(c.value, datetime):
                c.number_format = "DD/MM/YYYY"
    for col, larg in zip("ABCDEF", (12, 40, 12, 18, 18, 12)):
        ws.column_dimensions[col].width = larg
    wb.create_sheet("Planilha1")
    saida = PASTA / "nutricao_teste.xlsx"
    wb.save(saida)
    print(f"{saida}  ({len(LINHAS)} pacientes)")


if __name__ == "__main__":
    main()
